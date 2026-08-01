"""
╔══════════════════════════════════════════════════════════════╗
║  SCRAPER LOCALES/OFICINAS EN ALQUILER — AMPOSTA             ║
║  Criterios: 200-300 m² | Alquiler                           ║
║  v4.0 — Motor Playwright (renderiza JavaScript)             ║
╚══════════════════════════════════════════════════════════════╝

INSTALACIÓN (solo la primera vez):
    pip install playwright beautifulsoup4 pandas openpyxl lxml requests
    playwright install chromium

EJECUCIÓN:
    python scraper_oficinas_amposta.py

NOTA: Abre Chrome en modo invisible. No toques el ratón mientras corre.
"""

from playwright.sync_api import sync_playwright, TimeoutError as PWTimeout
from bs4 import BeautifulSoup
import requests, pandas as pd, time, random, re, logging
from datetime import datetime

# ──────────────────────────────────────────
#  CONFIGURACIÓN
# ──────────────────────────────────────────
MIN_M2          = 200
MAX_M2          = 300
INCLUIR_SIN_M2  = True     # True = incluye sin m² para revisión manual
HEADLESS        = True     # False = ver el navegador en pantalla (útil para depurar)
PAUSA_MIN       = 3        # segundos entre páginas
PAUSA_MAX       = 7

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s [%(levelname)s] %(message)s",
    handlers=[logging.StreamHandler(),
              logging.FileHandler("scraper_log.txt", encoding="utf-8")]
)
log = logging.getLogger(__name__)

UA = ("Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) "
      "AppleWebKit/537.36 (KHTML, like Gecko) "
      "Chrome/124.0.0.0 Safari/537.36")

# ──────────────────────────────────────────
#  UTILIDADES
# ──────────────────────────────────────────
def pausa(a=None, b=None):
    time.sleep(random.uniform(a or PAUSA_MIN, b or PAUSA_MAX))

def extraer_m2(txt):
    if not txt:
        return None
    m = re.search(r"\b(\d{2,4})\s*m[2²]", str(txt), re.IGNORECASE)
    return int(m.group(1)) if m else None

def ok_m2(m2):
    return INCLUIR_SIN_M2 if m2 is None else MIN_M2 <= m2 <= MAX_M2

def mk(portal, tit, prec, m2, det, url):
    return {"Portal": portal, "Título": (tit or "")[:200],
            "Precio": (prec or "")[:80], "Metros²": m2,
            "Detalles": (det or "")[:500], "URL": url or "",
            "Ciudad": "Amposta", "Operación": "Alquiler"}

def abs_url(href, base):
    if not href: return ""
    return href if href.startswith("http") else "/".join(base.split("/")[:3]) + "/" + href.lstrip("/")

# ──────────────────────────────────────────
#  MOTOR PLAYWRIGHT
# ──────────────────────────────────────────
_browser = None
_pw      = None

def init_browser():
    global _browser, _pw
    _pw      = sync_playwright().start()
    _browser = _pw.chromium.launch(
        headless=HEADLESS,
        args=["--no-sandbox", "--disable-dev-shm-usage",
              "--disable-blink-features=AutomationControlled"]
    )
    log.info("Navegador iniciado")

def close_browser():
    global _browser, _pw
    try:
        _browser.close(); _pw.stop()
    except Exception:
        pass

def nueva_pagina():
    ctx = _browser.new_context(
        user_agent=UA,
        locale="es-ES",
        viewport={"width": 1280, "height": 800},
        extra_http_headers={"Accept-Language": "es-ES,es;q=0.9,ca;q=0.8"}
    )
    # Ocultar rastro de automatización
    ctx.add_init_script("""
        Object.defineProperty(navigator, 'webdriver', {get: () => undefined});
        Object.defineProperty(navigator, 'plugins', {get: () => [1,2,3]});
        window.chrome = {runtime: {}};
    """)
    return ctx.new_page()

def get_html_pw(url, espera_selector=None, timeout=20000):
    """
    Abre url con el navegador, espera a que cargue el contenido y
    devuelve el HTML renderizado. Retorna None si falla.
    """
    page = nueva_pagina()
    try:
        page.goto(url, timeout=timeout, wait_until="domcontentloaded")
        if espera_selector:
            try:
                page.wait_for_selector(espera_selector, timeout=8000)
            except PWTimeout:
                pass   # si el selector no aparece, usamos lo que haya
        else:
            page.wait_for_timeout(3000)   # 3s para que cargue el JS
        return page.content()
    except PWTimeout:
        log.warning(f"Timeout en {url}")
        return None
    except Exception as e:
        log.error(f"Error Playwright en {url}: {e}")
        return None
    finally:
        try:
            page.context.close()
        except Exception:
            pass

def get_html_requests(url):
    """Para portales simples sin JS (pisos.com)."""
    hdrs = {"User-Agent": UA, "Accept-Language": "es-ES,es;q=0.9",
            "Accept": "text/html,application/xhtml+xml,*/*;q=0.8"}
    try:
        r = requests.get(url, headers=hdrs, timeout=15, allow_redirects=True)
        r.raise_for_status()
        return r.text
    except Exception as e:
        log.warning(f"requests error {url}: {e}")
        return None


# ══════════════════════════════════════════
#  PORTALES GRANDES
# ══════════════════════════════════════════

def scrape_idealista():
    portal  = "Idealista"
    results = []
    for page in range(1, 6):
        url = "https://www.idealista.com/alquiler-locales/amposta-tarragona/"
        if page > 1:
            url = f"https://www.idealista.com/alquiler-locales/amposta-tarragona/pagina-{page}.htm"
        log.info(f"[{portal}] pág.{page}")
        html = get_html_pw(url, espera_selector="article.item", timeout=25000)
        pausa()
        if not html:
            break
        if "captcha" in html.lower() or "robot" in html.lower():
            log.warning(f"[{portal}] ⚠ Captcha detectado")
            break
        soup = BeautifulSoup(html, "lxml")
        arts = soup.find_all("article", class_="item")
        log.info(f"[{portal}] {len(arts)} artículos encontrados")
        if not arts:
            break
        for art in arts:
            try:
                a    = art.find("a", class_="item-link")
                tit  = a.get_text(strip=True) if a else ""
                link = "https://www.idealista.com" + a["href"] if a and a.get("href") else ""
                prec = art.find("span", class_="item-price")
                prec = prec.get_text(strip=True) if prec else ""
                det  = art.find("div", class_="item-detail-char")
                det  = det.get_text(" ", strip=True) if det else ""
                m2   = extraer_m2(det) or extraer_m2(tit)
                if ok_m2(m2):
                    results.append(mk(portal, tit, prec, m2, det, link))
            except Exception as e:
                log.debug(f"[{portal}] {e}")
        if not soup.find("a", class_="icon-arrow-right-after"):
            break
    log.info(f"[{portal}] ✔ {len(results)}")
    return results


def scrape_fotocasa():
    portal  = "Fotocasa"
    results = []
    for page in range(1, 6):
        url = "https://www.fotocasa.es/es/alquiler/locales/amposta/todas-las-zonas/l"
        if page > 1:
            url += f"?page={page}"
        log.info(f"[{portal}] pág.{page}")
        # Esperar a que carguen las tarjetas
        html = get_html_pw(url, espera_selector="[class*='re-Card'],[class*='PropertyCard'],article", timeout=25000)
        pausa()
        if not html:
            break
        soup = BeautifulSoup(html, "lxml")
        items = (soup.find_all("article") or
                 soup.find_all("div", class_=re.compile(r"re-Card|PropertyCard|re-Searc", re.I)) or
                 soup.find_all("li",  class_=re.compile(r"re-Card", re.I)))
        log.info(f"[{portal}] {len(items)} items")
        if not items:
            break
        for item in items:
            try:
                a    = item.find("a", href=True)
                link = ("https://www.fotocasa.es" + a["href"]) if a else ""
                tit  = a.get_text(strip=True) if a else ""
                prec = item.find(class_=re.compile(r"price|precio", re.I))
                prec = prec.get_text(strip=True) if prec else ""
                txt  = item.get_text(" ", strip=True)
                m2   = extraer_m2(txt)
                if ok_m2(m2) and len(txt) > 30:
                    results.append(mk(portal, tit, prec, m2, txt, link))
            except Exception as e:
                log.debug(f"[{portal}] {e}")
        if not soup.find("a", {"aria-label": re.compile(r"siguiente|next", re.I)}):
            break
    log.info(f"[{portal}] ✔ {len(results)}")
    return results


def scrape_habitaclia():
    portal  = "Habitaclia"
    results = []
    for page in range(1, 6):
        url = "https://www.habitaclia.com/alquiler-locales_comerciales-amposta.htm"
        if page > 1:
            url = f"https://www.habitaclia.com/alquiler-locales_comerciales-amposta-{page}.htm"
        log.info(f"[{portal}] pág.{page}")
        html = get_html_pw(url, timeout=20000)
        pausa()
        if not html:
            break
        if any(x in html.lower() for x in ["no s'han trobat","no se han encontrado"]):
            break
        soup = BeautifulSoup(html, "lxml")
        items = (soup.find_all("article", class_=re.compile(r"js-list-item|list-item", re.I)) or
                 soup.find_all("li",      class_=re.compile(r"list-item", re.I)) or
                 soup.find_all("div",     class_=re.compile(r"property-info|js-property", re.I)))
        log.info(f"[{portal}] {len(items)} items")
        if not items:
            break
        for item in items:
            try:
                a    = item.find("a", href=True)
                link = abs_url(a["href"] if a else "", "https://www.habitaclia.com")
                tit  = item.find(["h2","h3"])
                tit  = tit.get_text(strip=True) if tit else ""
                prec = item.find(class_=re.compile(r"price|preu|precio", re.I))
                prec = prec.get_text(strip=True) if prec else ""
                txt  = item.get_text(" ", strip=True)
                m2   = extraer_m2(txt)
                if ok_m2(m2):
                    results.append(mk(portal, tit, prec, m2, txt, link))
            except Exception as e:
                log.debug(f"[{portal}] {e}")
        if not soup.find("a", rel="next"):
            break
    log.info(f"[{portal}] ✔ {len(results)}")
    return results


def scrape_pisos():
    # Pisos.com funciona con requests (HTML estático)
    portal  = "Pisos.com"
    results = []
    hdrs = {"User-Agent": UA, "Accept-Language": "es-ES,es;q=0.9"}
    for page in range(1, 6):
        url = "https://www.pisos.com/alquiler/locales-amposta/"
        if page > 1:
            url = f"https://www.pisos.com/alquiler/locales-amposta/{page}/"
        log.info(f"[{portal}] pág.{page}")
        try:
            r = requests.get(url, headers=hdrs, timeout=15)
            r.raise_for_status()
            html = r.text
        except Exception as e:
            log.warning(f"[{portal}] {e}")
            break
        pausa(2, 4)
        soup = BeautifulSoup(html, "lxml")
        items = soup.find_all("div", class_=re.compile(r"ad-preview|c-ad", re.I))
        log.info(f"[{portal}] {len(items)} items")
        if not items:
            break
        for item in items:
            try:
                a    = item.find("a", href=True)
                link = abs_url(a["href"] if a else "", "https://www.pisos.com")
                tit  = item.find(class_=re.compile(r"title|titulo", re.I))
                tit  = tit.get_text(strip=True) if tit else ""
                prec = item.find(class_=re.compile(r"price|precio", re.I))
                prec = prec.get_text(strip=True) if prec else ""
                txt  = item.get_text(" ", strip=True)
                m2   = extraer_m2(txt)
                if ok_m2(m2):
                    results.append(mk(portal, tit, prec, m2, txt, link))
            except Exception as e:
                log.debug(f"[{portal}] {e}")
        if not soup.find("a", rel="next"):
            break
    log.info(f"[{portal}] ✔ {len(results)}")
    return results


def scrape_yaencontre():
    portal  = "Yaencontre"
    results = []
    for page in range(1, 6):
        url = "https://www.yaencontre.com/alquiler/locales/amposta"
        if page > 1:
            url += f"?page={page}"
        log.info(f"[{portal}] pág.{page}")
        html = get_html_pw(url, espera_selector="[class*='card'],[class*='Card'],article", timeout=20000)
        pausa()
        if not html:
            break
        soup = BeautifulSoup(html, "lxml")
        items = (soup.find_all("div", class_=re.compile(r"card-property|CardProperty|re-Card|SearchResult", re.I)) or
                 soup.find_all("article") or
                 soup.find_all("li", class_=re.compile(r"item|result", re.I)))
        log.info(f"[{portal}] {len(items)} items")
        if not items:
            break
        for item in items:
            try:
                a    = item.find("a", href=True)
                link = abs_url(a["href"] if a else "", "https://www.yaencontre.com")
                tit  = item.find(["h2","h3"])
                tit  = tit.get_text(strip=True) if tit else ""
                prec = item.find(class_=re.compile(r"price|precio", re.I))
                prec = prec.get_text(strip=True) if prec else ""
                txt  = item.get_text(" ", strip=True)
                m2   = extraer_m2(txt)
                if ok_m2(m2) and len(txt) > 30:
                    results.append(mk(portal, tit, prec, m2, txt, link))
            except Exception as e:
                log.debug(f"[{portal}] {e}")
        if not soup.find("a", rel="next"):
            break
    log.info(f"[{portal}] ✔ {len(results)}")
    return results


def scrape_wallapop():
    portal  = "Wallapop"
    results = []
    for url in ["https://es.wallapop.com/inmobiliaria/alquilar-locales/amposta",
                "https://es.wallapop.com/inmobiliaria/alquilar-oficinas/amposta"]:
        log.info(f"[{portal}] {url}")
        html = get_html_pw(url, timeout=20000)
        pausa()
        if not html:
            continue
        soup = BeautifulSoup(html, "lxml")
        items = (soup.find_all("a",   class_=re.compile(r"ItemCard|item-card", re.I)) or
                 soup.find_all("div", class_=re.compile(r"ItemCard|listing-card|Card", re.I)) or
                 soup.find_all("article"))
        for item in items:
            try:
                a    = item if item.name == "a" else item.find("a", href=True)
                link = abs_url(a["href"] if a and a.get("href") else "", "https://es.wallapop.com")
                tit  = item.find(["h2","h3","p"])
                tit  = tit.get_text(strip=True) if tit else ""
                txt  = item.get_text(" ", strip=True)
                pm   = re.search(r"[\d.,]+\s*€", txt)
                prec = pm.group(0) if pm else ""
                m2   = extraer_m2(txt)
                if ok_m2(m2) and tit and len(txt) > 20:
                    results.append(mk(portal, tit, prec, m2, txt, link))
            except Exception as e:
                log.debug(f"[{portal}] {e}")
    log.info(f"[{portal}] ✔ {len(results)}")
    return results


# ══════════════════════════════════════════
#  INMOBILIARIAS LOCALES  (con Playwright)
#
#  Estas agencias usan JS. Con Playwright
#  podemos ver el HTML real. Usamos una
#  estrategia de descubrimiento: cargamos
#  la home y buscamos los enlaces internos.
# ══════════════════════════════════════════

KWORDS_ALQUILER = ["alquiler","lloguer","arrendament","rent"]
KWORDS_LOCAL    = ["local","oficina","comercial","negoci"]
KWORDS_ANUNCIO  = ["€","m²","m2","alquiler","lloguer","local","oficina","metros","metres","precio","preu"]

def _extraer_items_de_soup(soup, portal, base_url):
    results = []
    items = (soup.find_all("article") or
             soup.find_all("div",  class_=re.compile(r"property|inmueble|listing|card|propiedad|item|result", re.I)) or
             soup.find_all("li",   class_=re.compile(r"property|inmueble|listing|item|result", re.I)))
    for item in items:
        try:
            txt = item.get_text(" ", strip=True)
            if len(txt) < 40 or not any(k in txt.lower() for k in KWORDS_ANUNCIO):
                continue
            m2   = extraer_m2(txt)
            a    = item.find("a", href=True)
            link = abs_url(a["href"] if a else "", base_url)
            tit  = item.find(["h2","h3","h4"])
            tit  = tit.get_text(strip=True) if tit else txt[:80]
            prec = item.find(class_=re.compile(r"price|preu|precio", re.I))
            prec = prec.get_text(strip=True) if prec else ""
            if not prec:
                pm = re.search(r"[\d.,]+\s*€", txt)
                prec = pm.group(0) if pm else ""
            if ok_m2(m2) and tit and len(tit) > 5:
                results.append(mk(portal, tit, prec, m2, txt, link))
        except Exception:
            continue
    return results

def _descubrir_y_scrapear(nombre, home, urls_directas=None):
    """
    1. Prueba URLs directas con Playwright.
    2. Si no encuentra nada, carga la home con Playwright
       y descubre los enlaces internos de listado.
    3. Prueba esos enlaces.
    """
    results     = []
    probadas    = set()
    candidatos  = list(urls_directas or [])

    # Fase 1: URLs directas
    for url in candidatos:
        if url in probadas:
            continue
        probadas.add(url)
        log.info(f"[{nombre}] → {url}")
        html = get_html_pw(url, timeout=18000)
        pausa(2, 4)
        if not html:
            continue
        soup   = BeautifulSoup(html, "lxml")
        enc    = _extraer_items_de_soup(soup, nombre, url)
        log.info(f"[{nombre}]   {len(enc)} anuncios")
        results.extend(enc)
        if enc:
            log.info(f"[{nombre}] ✔ {len(results)} (URL directa)")
            return results

    # Fase 2: Descubrimiento desde home con Playwright
    if home not in probadas:
        log.info(f"[{nombre}] Descubriendo desde home (JS): {home}")
        html = get_html_pw(home, timeout=18000)
        pausa(2, 4)
        if html:
            soup = BeautifulSoup(html, "lxml")
            for a in soup.find_all("a", href=True):
                href  = a["href"].strip()
                texto = a.get_text(strip=True).lower()
                href_l = href.lower()
                if not href or href.startswith("#") or "javascript" in href_l:
                    continue
                if any(k in href_l or k in texto for k in KWORDS_ALQUILER + KWORDS_LOCAL):
                    full = abs_url(href, home)
                    if full and full not in probadas and "/".join(home.split("/")[:3]) in full:
                        candidatos.append(full)
            log.info(f"[{nombre}] {len(candidatos) - len(probadas)} candidatos descubiertos")

    # Fase 3: Probar candidatos descubiertos
    for url in candidatos:
        if url in probadas:
            continue
        probadas.add(url)
        log.info(f"[{nombre}] → {url}")
        html = get_html_pw(url, timeout=18000)
        pausa(2, 4)
        if not html:
            continue
        soup = BeautifulSoup(html, "lxml")
        enc  = _extraer_items_de_soup(soup, nombre, url)
        log.info(f"[{nombre}]   {len(enc)} anuncios")
        results.extend(enc)

    log.info(f"[{nombre}] ✔ Total: {len(results)}")
    return results


# — Cada agencia con sus URLs más probables —

def scrape_finques_marivent():
    return _descubrir_y_scrapear("Finques Marivent", "https://www.finquesmarivent.com", [
        "https://www.finquesmarivent.com/ca/propietats/",
        "https://www.finquesmarivent.com/es/propiedades/",
    ])

def scrape_finques_farnos():
    return _descubrir_y_scrapear("Finques Farnós", "https://www.finquesfarnos.com", [
        "https://www.finquesfarnos.com/propiedades/",
        "https://www.finquesfarnos.com/inmuebles/",
    ])

def scrape_romeu_immobles():
    return _descubrir_y_scrapear("Romeu Immobles", "https://www.romeuimmobles.com", [
        "https://www.romeuimmobles.com/localizacion/catalunya/tarragona/terres-del-ebre-delta-del-ebre/amposta/",
        "https://www.romeuimmobles.com/propiedades/?operacion=alquiler",
    ])

def scrape_trading_habitat():
    # ✅ URL verificada desde Google
    return _descubrir_y_scrapear("Trading Habitat", "https://www.tradinghabitat.com", [
        "https://www.tradinghabitat.com/alquileres/",
        "https://www.tradinghabitat.com/propiedades/",
    ])

def scrape_finques_garcia():
    return _descubrir_y_scrapear("Finques Garcia", "http://www.finquesgarcia.com", [
        "http://www.finquesgarcia.com/propiedades/",
        "http://www.finquesgarcia.com/inmuebles/",
    ])

def scrape_minerva():
    return _descubrir_y_scrapear("Minerva Inmobiliaria", "https://www.minervainmobiliaria.es", [
        "https://www.minervainmobiliaria.es/propiedades/",
        "https://www.minervainmobiliaria.es/alquiler/",
        "https://www.minervainmobiliaria.es/inmuebles/",
    ])

def scrape_homein():
    return _descubrir_y_scrapear("Home In", "http://www.homein.cat", [
        "http://www.homein.cat/ca/lloguer/",
        "http://www.homein.cat/es/alquiler/",
        "http://www.homein.cat/propietats/",
    ])

def scrape_ebre_taxacions():
    return _descubrir_y_scrapear("Ebre Taxacions", "http://www.ebretaxacions.com", [
        "http://www.ebretaxacions.com/propiedades/",
        "http://www.ebretaxacions.com/alquiler/",
    ])

def scrape_immoebre():
    return _descubrir_y_scrapear("ImmoEbre", "https://www.immoebre.com", [
        "https://www.immoebre.com/cerca/",
        "https://www.immoebre.com/buscar/",
        "https://www.immoebre.com/propietats/",
    ])

def scrape_delta_ebro():
    return _descubrir_y_scrapear("Delta del Ebro", "https://inmobiliaria-deltadelebro.com", [
        "https://inmobiliaria-deltadelebro.com/es/alquiler/",
        "https://inmobiliaria-deltadelebro.com/propiedades/",
    ])

def scrape_finques_roca():
    return _descubrir_y_scrapear("Finques Roca", "https://www.finquesroca.com", [
        "https://www.finquesroca.com/ca/lloguer/",
        "https://www.finquesroca.com/propietats/",
    ])

def scrape_iad():
    return _descubrir_y_scrapear("IAD España", "https://www.iadespana.es", [
        "https://www.iadespana.es/anuncios/alquiler/locales/amposta/",
        "https://www.iadespana.es/buscar/?operacion=alquiler&ciudad=amposta",
    ])

def scrape_safti():
    return _descubrir_y_scrapear("SAFTI España", "https://www.safti.es", [
        "https://www.safti.es/anuncios/?operacion=alquiler&localidad=amposta",
        "https://www.safti.es/propiedades/",
    ])


# ══════════════════════════════════════════
#  EXPORTAR EXCEL
# ══════════════════════════════════════════
COLORES = {"Idealista":"FFF3CD","Fotocasa":"FFE4CC","Habitaclia":"CCE5FF",
           "Pisos.com":"D4EDDA","Yaencontre":"E8D5F5","Wallapop":"F8D7E3"}

def exportar_excel(df, path):
    from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Resultados")
        wb = writer.book
        ws = writer.sheets["Resultados"]
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="2C3E50")
            cell.font = Font(bold=True, color="FFFFFF", size=11)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 28
        ws.auto_filter.ref = ws.dimensions
        ws.freeze_panes   = "A2"
        thin  = Side(style="thin", color="CCCCCC")
        borde = Border(left=thin, right=thin, top=thin, bottom=thin)
        cols  = [c.value for c in ws[1]]
        ip    = cols.index("Portal") + 1 if "Portal" in cols else None
        iu    = cols.index("URL")    + 1 if "URL"    in cols else None
        for r in range(2, ws.max_row + 1):
            p    = ws.cell(r, ip).value if ip else ""
            fill = PatternFill("solid", fgColor=COLORES.get(p, "F5F5F5"))
            for c in ws[r]:
                c.fill = fill; c.border = borde
                c.alignment = Alignment(wrap_text=True, vertical="top")
            if iu:
                uc = ws.cell(r, iu)
                if uc.value and str(uc.value).startswith("http"):
                    uc.hyperlink = uc.value
                    uc.font = Font(color="0563C1", underline="single")
        for col in ws.columns:
            w = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(w + 3, 70)
        ws2 = wb.create_sheet("Resumen")
        ws2.append(["Portal", "Anuncios"])
        ws2["A1"].font = Font(bold=True); ws2["B1"].font = Font(bold=True)
        for p, g in df.groupby("Portal"):
            ws2.append([p, len(g)])
    log.info(f"✅ Excel: {path}")


# ══════════════════════════════════════════
#  MAIN
# ══════════════════════════════════════════
SCRAPERS = [
    ("Idealista",            scrape_idealista),
    ("Fotocasa",             scrape_fotocasa),
    ("Habitaclia",           scrape_habitaclia),
    ("Pisos.com",            scrape_pisos),         # usa requests (funciona bien)
    ("Yaencontre",           scrape_yaencontre),
    ("Wallapop",             scrape_wallapop),
    ("Finques Marivent",     scrape_finques_marivent),
    ("Finques Farnós",       scrape_finques_farnos),
    ("Romeu Immobles",       scrape_romeu_immobles),
    ("Trading Habitat",      scrape_trading_habitat),
    ("Finques Garcia",       scrape_finques_garcia),
    ("Minerva Inmobiliaria", scrape_minerva),
    ("Home In",              scrape_homein),
    ("Ebre Taxacions",       scrape_ebre_taxacions),
    ("ImmoEbre",             scrape_immoebre),
    ("Delta del Ebro",       scrape_delta_ebro),
    ("Finques Roca",         scrape_finques_roca),
    ("IAD España",           scrape_iad),
    ("SAFTI España",         scrape_safti),
    # Ferrer International eliminado: dominio DNS inactivo
]

def main():
    log.info("=" * 60)
    log.info("  SCRAPER LOCALES/OFICINAS — AMPOSTA  v4.0 (Playwright)")
    log.info(f"  Filtro: {MIN_M2}–{MAX_M2} m²  |  Fuentes: {len(SCRAPERS)}")
    log.info("=" * 60)

    init_browser()
    todos   = []
    resumen = []

    try:
        for nombre, scraper in SCRAPERS:
            log.info(f"\n▶ {nombre}")
            try:
                res = scraper()
            except Exception as e:
                log.error(f"  Error en {nombre}: {e}")
                res = []
            todos.extend(res)
            resumen.append({"Portal": nombre, "Anuncios": len(res)})
            pausa()
    finally:
        close_browser()

    # Deduplicar por URL
    vistos, unicos = set(), []
    for r in todos:
        k = r.get("URL") or r.get("Título","")
        if k not in vistos:
            vistos.add(k); unicos.append(r)

    log.info("\n" + "=" * 60)
    log.info("  RESUMEN FINAL")
    for s in resumen:
        log.info(f"  {'✔' if s['Anuncios'] else '—'} {s['Portal']:<28} {s['Anuncios']:>4}")
    log.info(f"\n  TOTAL ÚNICO: {len(unicos)}")
    log.info("=" * 60)

    if not unicos:
        log.warning("Sin resultados. Prueba con HEADLESS = False para ver qué pasa en el navegador.")
        unicos = [{"Portal":"—","Título":"Sin resultados","Precio":"","Metros²":"",
                   "Detalles":"Ver log para detalles","URL":"","Ciudad":"Amposta","Operación":"Alquiler"}]

    df = pd.DataFrame(unicos)
    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    exportar_excel(df, f"oficinas_amposta_{ts}.xlsx")
    df.to_csv(f"oficinas_amposta_{ts}.csv", index=False, encoding="utf-8-sig")
    log.info(f"✅ CSV guardado")
    return df

if __name__ == "__main__":
    main()
